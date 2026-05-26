import csv
import sqlite3
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch

from openpyxl import Workbook

from portfolio_news_agent.portfolio_importer import (
    calculate_file_hash,
    import_portfolio_file,
    is_valid_symbol,
)
from portfolio_news_agent.storage import migrate


class PortfolioImporterTests(unittest.TestCase):
    def test_csv_import_saves_assets_and_source_hash(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            portfolio_path = Path(tmp_dir) / "portfolio.csv"
            self._write_csv(
                portfolio_path,
                [
                    {
                        "Symbol": "AEM",
                        "Name": "Agnico Eagle Mines",
                        "Price": "67.25",
                        "Quant Rating": "Buy",
                        "SA Analyst Ratings": "Buy",
                        "Wall Street Ratings": "Strong Buy",
                        "Sector": "Materials",
                        "Industry": "Gold",
                        "Asset Type": "Stock",
                    },
                    {
                        "Symbol": "QQQ",
                        "Name": "Invesco QQQ Trust",
                        "Price": "450.10",
                        "Quant Rating": "Hold",
                        "SA Analyst Ratings": "",
                        "Wall Street Ratings": "",
                        "Sector": "ETF",
                        "Industry": "Large Cap Growth",
                        "Asset Type": "ETF",
                    },
                ],
            )
            connection = sqlite3.connect(":memory:")
            migrate(connection)

            result = import_portfolio_file(connection, portfolio_path)
            assets = self._assets(connection, result.import_id)
            import_row = connection.execute(
                "SELECT source_path, source_hash FROM portfolio_imports WHERE id = ?",
                (result.import_id,),
            ).fetchone()
            expected_hash = calculate_file_hash(portfolio_path)

        self.assertEqual(result.assets_imported, 2)
        self.assertEqual(import_row["source_hash"], expected_hash)
        self.assertEqual([asset["symbol"] for asset in assets], ["AEM", "QQQ"])
        self.assertEqual(assets[0]["quant_rating"], "Buy")
        self.assertEqual(assets[0]["sa_analyst_rating"], "Buy")
        self.assertEqual(assets[0]["wall_street_rating"], "Strong Buy")
        self.assertEqual(assets[0]["sector"], "Materials")
        self.assertEqual(assets[0]["industry"], "Gold")
        self.assertEqual(assets[0]["asset_type"], "Stock")

    def test_invalid_symbol_filtering(self):
        invalid_values = ["", "123", "2026-05-23", "not a symbol", "TOO-LONG-SYMBOL"]
        for value in invalid_values:
            with self.subTest(value=value):
                self.assertFalse(is_valid_symbol(value))

        for value in ["AEM", "ASML", "QQQ", "SILJ", "BRK.B"]:
            with self.subTest(value=value):
                self.assertTrue(is_valid_symbol(value))

    def test_csv_import_ignores_non_symbol_rows(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            portfolio_path = Path(tmp_dir) / "portfolio.csv"
            self._write_csv(
                portfolio_path,
                [
                    {"Symbol": "AEM", "Name": "Agnico Eagle Mines"},
                    {"Symbol": "123", "Name": "Lot row"},
                    {"Symbol": "2026-05-23", "Name": "Date row"},
                    {"Symbol": "", "Name": "Blank row"},
                    {"Symbol": "BRK.B", "Name": "Berkshire Hathaway"},
                ],
            )
            connection = sqlite3.connect(":memory:")
            migrate(connection)

            result = import_portfolio_file(connection, portfolio_path)
            symbols = [asset["symbol"] for asset in self._assets(connection, result.import_id)]

        self.assertEqual(result.assets_imported, 2)
        self.assertEqual(symbols, ["AEM", "BRK.B"])

    def test_excel_import_falls_back_to_first_sheet_with_symbol_column(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "portfolio.xlsx"
            workbook = Workbook()
            workbook.active.title = "Notes"
            workbook.active.append(["Not Symbol", "Name"])
            holdings = workbook.create_sheet("Holdings")
            holdings.append(["Name", "Symbol", "Price"])
            holdings.append(["Agnico Eagle Mines", "AEM", 67.25])
            holdings.append(["Berkshire Hathaway", "BRK.B", 410.0])
            workbook.save(workbook_path)
            connection = sqlite3.connect(":memory:")
            migrate(connection)

            result = import_portfolio_file(connection, workbook_path)
            assets = self._assets(connection, result.import_id)

        self.assertEqual(result.source_sheet, "Holdings")
        self.assertEqual([asset["symbol"] for asset in assets], ["AEM", "BRK.B"])

    def test_excel_import_prefers_clean_seeking_alpha_sheet(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "portfolio.xlsx"
            workbook = Workbook()
            workbook.active.title = "Holdings"
            workbook.active.append(["Symbol", "Name"])
            workbook.active.append(["BAD", "Wrong sheet"])
            clean = workbook.create_sheet("Excel view including prices")
            clean.append(["Symbol", "Name", "Price"])
            clean.append(["SILJ", "ETFMG Prime Junior Silver Miners ETF", 10.5])
            workbook.save(workbook_path)
            connection = sqlite3.connect(":memory:")
            migrate(connection)

            result = import_portfolio_file(connection, workbook_path)
            assets = self._assets(connection, result.import_id)

        self.assertEqual(result.source_sheet, "Excel view including prices")
        self.assertEqual([asset["symbol"] for asset in assets], ["SILJ"])

    def test_changed_portfolio_hash_creates_distinct_import(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            portfolio_path = Path(tmp_dir) / "portfolio.csv"
            self._write_csv(portfolio_path, [{"Symbol": "AEM", "Name": "Agnico Eagle Mines"}])
            connection = sqlite3.connect(":memory:")
            migrate(connection)

            first = import_portfolio_file(connection, portfolio_path)
            self._write_csv(portfolio_path, [{"Symbol": "NEM", "Name": "Newmont"}])
            second = import_portfolio_file(connection, portfolio_path)
            imports = connection.execute(
                "SELECT id, source_hash FROM portfolio_imports ORDER BY id"
            ).fetchall()

        self.assertNotEqual(first.import_id, second.import_id)
        self.assertEqual(len(imports), 2)
        self.assertNotEqual(imports[0]["source_hash"], imports[1]["source_hash"])

    def test_excel_import_falls_back_to_xml_values_when_openpyxl_rejects_formatting(self):
        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "portfolio.xlsx"
            self._write_minimal_xlsx(
                workbook_path,
                sheet_name="Excel view including prices",
                rows=[
                    ["Symbol", "Name", "Price"],
                    ["AEM", "Agnico Eagle Mines", "67.25"],
                    ["BRK.B", "Berkshire Hathaway", "410"],
                ],
            )
            connection = sqlite3.connect(":memory:")
            migrate(connection)

            with patch(
                "portfolio_news_agent.portfolio_importer.load_workbook",
                side_effect=ValueError("conditional formatting operator"),
            ):
                result = import_portfolio_file(connection, workbook_path)
            assets = self._assets(connection, result.import_id)

        self.assertEqual(result.source_sheet, "Excel view including prices")
        self.assertEqual([asset["symbol"] for asset in assets], ["AEM", "BRK.B"])
        self.assertEqual(assets[0]["price"], 67.25)

    def test_excel_import_falls_back_to_xml_values_when_openpyxl_row_iteration_fails(self):
        class BrokenWorkbook:
            sheetnames = ["Excel view including prices"]

            def __getitem__(self, name):
                return self

            def iter_rows(self, *args, **kwargs):
                raise ValueError("conditional formatting operator")

            def close(self):
                pass

        with tempfile.TemporaryDirectory() as tmp_dir:
            workbook_path = Path(tmp_dir) / "portfolio.xlsx"
            self._write_minimal_xlsx(
                workbook_path,
                sheet_name="Excel view including prices",
                rows=[["Symbol", "Name"], ["AEM", "Agnico Eagle Mines"]],
            )
            connection = sqlite3.connect(":memory:")
            migrate(connection)

            with patch(
                "portfolio_news_agent.portfolio_importer.load_workbook",
                return_value=BrokenWorkbook(),
            ):
                result = import_portfolio_file(connection, workbook_path)
            assets = self._assets(connection, result.import_id)

        self.assertEqual(result.source_sheet, "Excel view including prices")
        self.assertEqual([asset["symbol"] for asset in assets], ["AEM"])

    def _write_csv(self, path: Path, rows: list[dict[str, str]]) -> None:
        fieldnames = sorted({key for row in rows for key in row})
        with path.open("w", newline="", encoding="utf-8") as handle:
            writer = csv.DictWriter(handle, fieldnames=fieldnames)
            writer.writeheader()
            writer.writerows(rows)

    def _assets(self, connection: sqlite3.Connection, import_id: int) -> list[sqlite3.Row]:
        return connection.execute(
            "SELECT * FROM assets WHERE import_id = ? ORDER BY id",
            (import_id,),
        ).fetchall()

    def _write_minimal_xlsx(self, path: Path, sheet_name: str, rows: list[list[str]]) -> None:
        def cell_name(row_index: int, column_index: int) -> str:
            return f"{chr(64 + column_index)}{row_index}"

        sheet_rows = []
        for row_index, row in enumerate(rows, start=1):
            cells = []
            for column_index, value in enumerate(row, start=1):
                cells.append(
                    f'<c r="{cell_name(row_index, column_index)}" t="inlineStr">'
                    f"<is><t>{value}</t></is></c>"
                )
            sheet_rows.append(f'<row r="{row_index}">{"".join(cells)}</row>')

        with zipfile.ZipFile(path, "w") as archive:
            archive.writestr(
                "[Content_Types].xml",
                """<?xml version="1.0" encoding="UTF-8"?>
<Types xmlns="http://schemas.openxmlformats.org/package/2006/content-types">
  <Default Extension="rels" ContentType="application/vnd.openxmlformats-package.relationships+xml"/>
  <Default Extension="xml" ContentType="application/xml"/>
  <Override PartName="/xl/workbook.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet.main+xml"/>
  <Override PartName="/xl/worksheets/sheet1.xml" ContentType="application/vnd.openxmlformats-officedocument.spreadsheetml.worksheet+xml"/>
</Types>""",
            )
            archive.writestr(
                "_rels/.rels",
                """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/officeDocument" Target="xl/workbook.xml"/>
</Relationships>""",
            )
            archive.writestr(
                "xl/workbook.xml",
                f"""<?xml version="1.0" encoding="UTF-8"?>
<workbook xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main" xmlns:r="http://schemas.openxmlformats.org/officeDocument/2006/relationships">
  <sheets><sheet name="{sheet_name}" sheetId="1" r:id="rId1"/></sheets>
</workbook>""",
            )
            archive.writestr(
                "xl/_rels/workbook.xml.rels",
                """<?xml version="1.0" encoding="UTF-8"?>
<Relationships xmlns="http://schemas.openxmlformats.org/package/2006/relationships">
  <Relationship Id="rId1" Type="http://schemas.openxmlformats.org/officeDocument/2006/relationships/worksheet" Target="worksheets/sheet1.xml"/>
</Relationships>""",
            )
            archive.writestr(
                "xl/worksheets/sheet1.xml",
                f"""<?xml version="1.0" encoding="UTF-8"?>
<worksheet xmlns="http://schemas.openxmlformats.org/spreadsheetml/2006/main">
  <sheetData>{"".join(sheet_rows)}</sheetData>
</worksheet>""",
            )
