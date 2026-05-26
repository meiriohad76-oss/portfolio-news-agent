from __future__ import annotations

import csv
import hashlib
import re
import sqlite3
import zipfile
from dataclasses import dataclass
from pathlib import Path
from typing import Any
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

from portfolio_news_agent.storage import create_portfolio_import, insert_asset


PREFERRED_EXCEL_SHEET = "Excel view including prices"
_SYMBOL_PATTERN = re.compile(r"^[A-Z]{1,6}(?:[.-][A-Z]{1,2})?$")


@dataclass(frozen=True)
class PortfolioImportResult:
    import_id: int
    source_path: Path
    source_hash: str
    assets_imported: int
    source_sheet: str | None = None


def import_portfolio_file(
    connection: sqlite3.Connection,
    portfolio_path: str | Path,
) -> PortfolioImportResult:
    path = Path(portfolio_path)
    if not path.exists():
        raise FileNotFoundError(path)

    source_hash = calculate_file_hash(path)
    source_sheet: str | None = None
    if path.suffix.lower() == ".csv":
        rows = _read_csv_rows(path)
    elif path.suffix.lower() in {".xlsx", ".xlsm"}:
        source_sheet, rows = _read_excel_rows(path)
    else:
        raise ValueError(f"Unsupported portfolio file type: {path.suffix}")

    import_id = create_portfolio_import(
        connection,
        source_path=str(path),
        source_hash=source_hash,
    )

    assets_imported = 0
    for row in rows:
        symbol = _normalize_symbol(row.get("symbol"))
        if not is_valid_symbol(symbol):
            continue
        insert_asset(
            connection,
            import_id=import_id,
            symbol=symbol,
            name=_clean_text(row.get("name")),
            asset_type=_clean_text(row.get("asset_type")),
            sector=_clean_text(row.get("sector")),
            industry=_clean_text(row.get("industry")),
            price=_parse_float(row.get("price")),
            quant_rating=_clean_text(row.get("quant_rating")),
            sa_analyst_rating=_clean_text(row.get("sa_analyst_rating")),
            wall_street_rating=_clean_text(row.get("wall_street_rating")),
        )
        assets_imported += 1

    return PortfolioImportResult(
        import_id=import_id,
        source_path=path,
        source_hash=source_hash,
        assets_imported=assets_imported,
        source_sheet=source_sheet,
    )


def calculate_file_hash(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def is_valid_symbol(value: Any) -> bool:
    symbol = _normalize_symbol(value)
    if not symbol:
        return False
    return bool(_SYMBOL_PATTERN.fullmatch(symbol))


def _read_csv_rows(path: Path) -> list[dict[str, Any]]:
    with path.open(newline="", encoding="utf-8-sig") as handle:
        reader = csv.DictReader(handle)
        return [_normalize_row(row) for row in reader]


def _read_excel_rows(path: Path) -> tuple[str, list[dict[str, Any]]]:
    try:
        workbook = load_workbook(path, data_only=True, read_only=True)
    except ValueError:
        return _read_excel_rows_from_xml(path)
    try:
        sheet = _select_sheet(workbook)
        try:
            rows = list(sheet.iter_rows(values_only=True))
        except ValueError:
            return _read_excel_rows_from_xml(path)
        if not rows:
            return sheet.title, []
        headers = [_normalize_header(value) for value in rows[0]]
        data_rows = []
        for values in rows[1:]:
            data_rows.append(
                {
                    headers[index]: value
                    for index, value in enumerate(values)
                    if index < len(headers) and headers[index]
                }
            )
        return sheet.title, [_normalize_row(row) for row in data_rows]
    finally:
        workbook.close()


def _read_excel_rows_from_xml(path: Path) -> tuple[str, list[dict[str, Any]]]:
    with zipfile.ZipFile(path) as archive:
        shared_strings = _read_shared_strings(archive)
        sheets = _read_workbook_sheets(archive)
        selected_name, selected_path = _select_xml_sheet(archive, sheets, shared_strings)
        rows = _read_sheet_rows(archive, selected_path, shared_strings)

    if not rows:
        return selected_name, []
    headers = [_normalize_header(value) for value in rows[0]]
    data_rows = []
    for values in rows[1:]:
        data_rows.append(
            {
                headers[index]: value
                for index, value in enumerate(values)
                if index < len(headers) and headers[index]
            }
        )
    return selected_name, [_normalize_row(row) for row in data_rows]


def _read_workbook_sheets(archive: zipfile.ZipFile) -> list[tuple[str, str]]:
    namespace = {
        "main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main",
        "rel": "http://schemas.openxmlformats.org/officeDocument/2006/relationships",
        "pkgrel": "http://schemas.openxmlformats.org/package/2006/relationships",
    }
    workbook_root = ET.fromstring(archive.read("xl/workbook.xml"))
    rels_root = ET.fromstring(archive.read("xl/_rels/workbook.xml.rels"))
    relationship_targets = {
        relationship.attrib["Id"]: relationship.attrib["Target"]
        for relationship in rels_root.findall("pkgrel:Relationship", namespace)
    }
    sheets = []
    for sheet in workbook_root.findall("main:sheets/main:sheet", namespace):
        relationship_id = sheet.attrib[
            "{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id"
        ]
        target = relationship_targets[relationship_id]
        sheets.append((sheet.attrib["name"], _normalize_xlsx_target(target)))
    return sheets


def _select_xml_sheet(
    archive: zipfile.ZipFile,
    sheets: list[tuple[str, str]],
    shared_strings: list[str],
) -> tuple[str, str]:
    sheet_by_name = {name: path for name, path in sheets}
    if PREFERRED_EXCEL_SHEET in sheet_by_name:
        return PREFERRED_EXCEL_SHEET, sheet_by_name[PREFERRED_EXCEL_SHEET]

    for sheet_name, sheet_path in sheets:
        rows = _read_sheet_rows(archive, sheet_path, shared_strings, max_rows=1)
        headers = {_normalize_header(value) for value in (rows[0] if rows else [])}
        if "symbol" in headers:
            return sheet_name, sheet_path
    raise ValueError("No worksheet with a Symbol column was found")


def _read_shared_strings(archive: zipfile.ZipFile) -> list[str]:
    if "xl/sharedStrings.xml" not in archive.namelist():
        return []
    namespace = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(archive.read("xl/sharedStrings.xml"))
    strings = []
    for item in root.findall("main:si", namespace):
        parts = [node.text or "" for node in item.findall(".//main:t", namespace)]
        strings.append("".join(parts))
    return strings


def _read_sheet_rows(
    archive: zipfile.ZipFile,
    sheet_path: str,
    shared_strings: list[str],
    max_rows: int | None = None,
) -> list[list[Any]]:
    namespace = {"main": "http://schemas.openxmlformats.org/spreadsheetml/2006/main"}
    root = ET.fromstring(archive.read(sheet_path))
    rows = []
    for row in root.findall(".//main:sheetData/main:row", namespace):
        values_by_index: dict[int, Any] = {}
        for cell in row.findall("main:c", namespace):
            index = _column_index(cell.attrib.get("r", ""))
            values_by_index[index] = _cell_value(cell, shared_strings, namespace)
        if values_by_index:
            row_values = [
                values_by_index.get(index)
                for index in range(1, max(values_by_index) + 1)
            ]
            rows.append(row_values)
        if max_rows is not None and len(rows) >= max_rows:
            break
    return rows


def _cell_value(
    cell: ET.Element,
    shared_strings: list[str],
    namespace: dict[str, str],
) -> Any:
    cell_type = cell.attrib.get("t")
    if cell_type == "inlineStr":
        return "".join(node.text or "" for node in cell.findall(".//main:t", namespace))

    value_node = cell.find("main:v", namespace)
    if value_node is None or value_node.text is None:
        return None
    if cell_type == "s":
        return shared_strings[int(value_node.text)]
    return value_node.text


def _column_index(reference: str) -> int:
    letters = "".join(character for character in reference if character.isalpha())
    index = 0
    for letter in letters.upper():
        index = index * 26 + (ord(letter) - ord("A") + 1)
    return index or 1


def _normalize_xlsx_target(target: str) -> str:
    if target.startswith("/"):
        return target.lstrip("/")
    if target.startswith("xl/"):
        return target
    return f"xl/{target}"


def _select_sheet(workbook: Any) -> Any:
    if PREFERRED_EXCEL_SHEET in workbook.sheetnames:
        return workbook[PREFERRED_EXCEL_SHEET]

    for sheet in workbook.worksheets:
        first_row = next(sheet.iter_rows(min_row=1, max_row=1, values_only=True), ())
        headers = {_normalize_header(value) for value in first_row}
        if "symbol" in headers:
            return sheet

    raise ValueError("No worksheet with a Symbol column was found")


def _normalize_row(row: dict[str, Any]) -> dict[str, Any]:
    normalized: dict[str, Any] = {}
    for key, value in row.items():
        normalized_key = _normalize_header(key)
        if normalized_key:
            normalized[normalized_key] = value
    return normalized


def _normalize_header(value: Any) -> str:
    text = _clean_text(value)
    if text is None:
        return ""
    key = re.sub(r"[^a-z0-9]+", "_", text.lower()).strip("_")
    aliases = {
        "sa_analyst_ratings": "sa_analyst_rating",
        "wall_street_ratings": "wall_street_rating",
        "quant_ratings": "quant_rating",
        "asset_type": "asset_type",
    }
    return aliases.get(key, key)


def _normalize_symbol(value: Any) -> str:
    text = _clean_text(value)
    if text is None:
        return ""
    return text.upper()


def _clean_text(value: Any) -> str | None:
    if value is None:
        return None
    text = str(value).strip()
    if not text:
        return None
    return text


def _parse_float(value: Any) -> float | None:
    text = _clean_text(value)
    if text is None:
        return None
    try:
        return float(text.replace(",", ""))
    except ValueError:
        return None
